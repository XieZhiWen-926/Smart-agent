# -*- coding: utf-8 -*-
"""校验 .docx：结构（样式/表格/分节）+ 内容覆盖度（与源 HTML 对比）。"""
import html as H
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

DOCX = Path(sys.argv[1])
SRC_HTML = Path(sys.argv[2])

d = Document(str(DOCX))

print("=== 1. 基本结构 ===")
print(f"  段落总数 {len(d.paragraphs)}｜表格 {len(d.tables)}｜分节 {len(d.sections)}")
sec = d.sections[0]
print(f"  页面 {sec.page_width.cm:.1f} x {sec.page_height.cm:.1f} cm")
print(f"  页边距 上{sec.top_margin.cm:.1f} 下{sec.bottom_margin.cm:.1f} "
      f"左{sec.left_margin.cm:.1f} 右{sec.right_margin.cm:.1f} cm")

print("=== 2. 标题样式分布 ===")
from collections import Counter

cnt = Counter(p.style.name for p in d.paragraphs if p.text.strip())
for k, v in cnt.most_common(12):
    print(f"  {k}: {v}")

print("=== 3. 表格 ===")
for i, t in enumerate(d.tables, 1):
    hdr = [c.text.strip()[:12] for c in t.rows[0].cells]
    print(f"  表{i}: {len(t.rows)}行 x {len(t.columns)}列  表头={hdr}")

print("=== 4. 内容覆盖度（源 HTML 可见文本 → docx 文本）===")
h = SRC_HTML.read_text(encoding="utf-8")
h_body = h.split('role="body"')[1] if 'role="body"' in h else h
src = H.unescape(re.sub(r"<[^>]+>", " ", h_body))
src_flat = re.sub(r"\s+", "", src)

docx_txt = "\n".join(p.text for p in d.paragraphs)
for t in d.tables:
    for r in t.rows:
        for c in r.cells:
            docx_txt += "\n" + c.text
docx_flat = re.sub(r"\s+", "", H.unescape(docx_txt))

# 按句号切句，逐句检查是否出现在 docx
sents = [s for s in re.split(r"[。；\n]", src_flat) if len(s) >= 12]
miss = [s for s in sents if s not in docx_flat]
print(f"  源句 {len(sents)} 条，未出现在 docx 的 {len(miss)} 条")
for m in miss[:12]:
    print("   MISS:", m[:80])

print("=== 5. 段落样例 ===")
for i, p in enumerate(d.paragraphs):
    if not p.text.strip():
        continue
    if p.style.name.startswith("Heading") or i < 4:
        print(f"  [{p.style.name}] {p.text[:70]}")
    if i > 60:
        break
